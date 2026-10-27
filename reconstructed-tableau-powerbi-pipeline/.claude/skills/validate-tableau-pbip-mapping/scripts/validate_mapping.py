#!/usr/bin/env python3
"""
Validate a real build-tableau-pbip output: for every Tableau field, report exactly what it
became in Power BI, and for everything that did not make it, why -- plus a suggested fix.

RESIDUE.md already says what did not translate. Nothing says what a translated field actually
BECAME -- that answer already exists, in build-tableau-pbip's own _work/name_authority.json (the
Binding Ledger: every emitted table/column/measure, keyed by its Tableau source identity) and
_work/ir.json (the source formulas). This script is mostly a FORMATTER over data that already
exists, not a new extractor -- it cross-references those two files plus the emitted TMDL (for the
actual DAX text) and _work/residue.json (the same data RESIDUE.md renders as prose, structured).

Per workbook, two outputs:
  <name>_mapping.xlsx  -- Summary + Measures + Columns + Parameters + Gaps + Unaccounted sheets
  VALIDATE.dax         -- one EVALUATE query per mapped measure/column, for Power BI Desktop's
                          DAX query view, plus a commented GAPS section (reason + suggested fix)
                          for everything that was declined

Suggested fixes are looked up from FIX_RULES, a table of (pattern -> fix) keyed on the build's
OWN decline-reason text, not re-derived per instance. If a reason matches no rule, the report
says so plainly (GENERIC_FIX) rather than fabricating advice -- the same "never guess" convention
inventory-tableau-assets uses for `absent?` and map-tableau-powerbi-translation uses for
UNAPPORTIONED.

Usage:
  python3 validate_mapping.py <build-tableau-pbip --out folder> --out <output-folder>
"""
import os
import re
import sys
import json
import textwrap
import argparse
import collections

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
except ImportError:
    openpyxl = None


# ============================================================ fix-rule table
# Ordered: first pattern that matches the reason text wins. Each fn receives the regex Match, so
# rules can pull specifics (a function name, a tier, a table count) out of the reason and put
# them back into the fix. Grounded on the reason strings dax.py/tmdl.py actually raise -- see
# SKILL.md for how this catalogue was built, so a new decline path can be given a rule the same
# way.
_DATETRUNC_UNIT = re.compile(r"DATETRUNC\(\s*['\"](\w+)['\"]", re.I)
_DATEPART_UNIT = re.compile(r"DATEPART\(\s*['\"](\w+)['\"]", re.I)

_DATETRUNC_PATTERNS = {
    "year": "DATE(YEAR(<date>), 1, 1)",
    "quarter": "DATE(YEAR(<date>), 3*QUARTER(<date>)-2, 1)",
    "month": "DATE(YEAR(<date>), MONTH(<date>), 1)",
    "week": "<date> - WEEKDAY(<date>, 1) + 1  (Sunday-start week; adjust the WEEKDAY type code "
            "if Tableau's week start differs)",
    "day": "<date> itself -- already day granularity",
}

_DATEPART_PATTERNS = {
    "year": "YEAR(<date>)",
    "quarter": "QUARTER(<date>)",
    "month": "MONTH(<date>)",
    "day": "DAY(<date>)",
    "dayofyear": "<date> - DATE(YEAR(<date>), 1, 1) + 1",
    "weekday": "WEEKDAY(<date>)",
    "week": "WEEKNUM(<date>) -- verify the week-start/ISO convention matches Tableau's",
    "hour": "HOUR(<date>)",
    "minute": "MINUTE(<date>)",
    "second": "SECOND(<date>)",
}


def _unit_fix(formula, unit_re, patterns, label):
    """DATETRUNC/DATEPART both decline with a fixed 'unit' PLACEHOLDER in the reason text --
    the real Tableau unit ('month', 'hour', ...) only exists in the formula, not the reason.
    If the formula parses cleanly, name that specific unit's DAX pattern; otherwise, since a
    reviewer still needs a starting point, list every unit this pipeline knows a pattern for."""
    found = unit_re.search(formula or "")
    unit = found.group(1).lower() if found else None
    if unit and unit in patterns:
        return (f"This call uses unit '{unit}' -> {patterns[unit]} (<date> is this formula's own "
                f"date argument).")
    listing = "; ".join(f"{u} -> {p}" for u, p in patterns.items())
    return f"{label} has no single DAX function; direct per-unit patterns: {listing}."


FIX_RULES = [
    (re.compile(r"^depends on '.*?', which was itself declined \(tier (\w+)\)"),
     lambda m: f"Translate the dependency first (tier {m.group(1)}) -- this calc recovers "
               f"automatically once that one does, no independent fix needed here."),

    (re.compile(r"table calculation -- addressing/partitioning is not in the formula text"),
     lambda m: "Open this field's \"Edit Table Calculation\" dialog in Tableau to read the real "
               "addressing/partitioning fields, then hand-author the DAX with a RANKX/WINDOW-"
               "style pattern using ALLEXCEPT on those partition fields."),

    (re.compile(r"level-of-detail expression"),
     lambda m: "Tableau INCLUDE/EXCLUDE LOD expressions compute at a view-adjusted granularity "
               "that doesn't exist at model-build time -- there's no formula-text-only "
               "translation. Hand-author the DAX using CALCULATE + ALLEXCEPT (for the FIXED-"
               "like partition an INCLUDE/EXCLUDE implies) or ALLSELECTED/KEEPFILTERS (to "
               "respect the view's own filter context), against the explicit dimension fields "
               "named in the LOD's bracket list, then wire the result in as a measure."),

    (re.compile(r"the two sides compute '.*?' differently"),
     lambda m: "This column was excluded from its pair's Blend Key because the two tables "
               "compute it with structurally different formulas (e.g. one side concatenates "
               "several fields, the other is a bare field alias) -- there is no DAX fix here. "
               "In Power BI Desktop, either add a calculated column on the mismatched side that "
               "reproduces the OTHER side's exact formula (so the two are genuinely the same "
               "value), or accept that this field cannot key the blend and check whether the "
               "relationship's remaining shared column(s) are still enough to relate the tables."),

    (re.compile(r"Tableau COUNTD\(<expression>\) needs a single table to project into "
                r"SELECTCOLUMNS"),
     lambda m: 'Rewrite as COUNTROWS(DISTINCT(SELECTCOLUMNS(<table>, "x", <expression>))) -- a '
               "calculated-table restructure, not a function swap. The expression spans more "
               "than one table, so SELECTCOLUMNS has no single table to project into; move the "
               "calc to one source table first (or build a bridge table) before this can "
               "translate."),

    (re.compile(r"Tableau (SUM|MIN|MAX|AVG)\(<expression>\) needs the DAX iterator (\w+X), "
                r"which requires ONE table to iterate; the expression references (\d+)"),
     lambda m: (f"The expression references 0 tables -- check the source formula for a "
                f"constant-only aggregate; that is likely a Tableau authoring artifact, not a "
                f"real calculation." if m.group(3) == "0" else
                f"Rewrite using {m.group(2)} over a single table."
                + ("" if m.group(3) in ("0", "1") else " If the fields genuinely span multiple "
                   "tables, they need a bridge/star-schema table before this can be an "
                   "iterator."))),

    (re.compile(r"aggregates a field Tableau itself declares datatype 'string'"),
     lambda m: "The field's Tableau datatype is 'string' even though its values are numeric "
               "text, so DAX's MIN/MAX/MEDIAN (which require a numeric/date column) were "
               "declined rather than guessing. Fix in Power Query: add a \"Changed Type\" step "
               "casting the column to Whole Number/Decimal (or Date) before it reaches the "
               "model, then re-run the build -- the cast column translates cleanly with no "
               "VALUE() workaround needed."),

    (re.compile(r"comparison against NULL"),
     lambda m: "Rewrite the comparison using ISBLANK(): DAX evaluates BLANK() = 0 as TRUE, so "
               "`[x] = NULL` is NOT equivalent to `ISBLANK([x])` whenever [x] can legitimately "
               "be 0 or an empty string. Pattern: "
               "IF(ISBLANK([x]), <null-branch>, <non-null-branch>) instead of a literal NULL "
               "substitution."),

    (re.compile(r"unresolved field reference \[tokn_(\w+)\] -- no field by that name is defined"),
     lambda m: (
        f"CHECK FOR A NAMING-MISMATCH BUG IN THE SOURCE TABLEAU WORKBOOK before treating this as "
        f"a normal rename/delete: a 'tokn_' prefix on an otherwise-ordinary field name "
        f"([tokn_{m.group(1)}]) is a known real-corpus signature of a calc formula that was "
        f"copy-pasted from a worksheet built against a DIFFERENT table (one that legitimately "
        f"has a 'tokn_'-prefixed column, e.g. a token/masked-value convention) onto THIS table, "
        f"without updating the field reference to match this table's own naming. VERIFY by "
        f"checking whether this table already has a column named exactly '[{m.group(1)}]' "
        f"(same suffix, no prefix) -- if it does, this is almost certainly the bug, and the fix "
        f"is FIX IN TABLEAU FIRST, not in Power BI: open the workbook, edit this calculated "
        f"field, and replace every '[tokn_{m.group(1)}]' with '[{m.group(1)}]', then re-run this "
        f"build (Tableau itself would already be showing this exact field broken in Desktop, "
        f"independent of any migration). If no matching unprefixed column exists on this table, "
        f"this is instead a genuine missing-field problem -- fall back to the general "
        f"rename/delete guidance below.")),

    (re.compile(r"unresolved field reference \[.*?\] -- no field by that name is defined"),
     lambda m: "The referenced field does not exist anywhere in the workbook -- check Tableau "
               "for a rename/delete. This is a source-data problem, not a translation gap; fix "
               "it in Tableau (or drop the dependent calc) before re-running the build."),

    (re.compile(r"unresolved parameter reference"),
     lambda m: "The referenced parameter does not exist -- check the Parameters datasource in "
               "Tableau for a rename/delete."),

    (re.compile(r"no DAX equivalent for Tableau function DATEPARSE\(\)"),
     lambda m: "DATEPARSE(format, string) has no single DAX function. Rewrite using DATEVALUE() "
               "if the string is already a recognizable date format, or a manual "
               "VALUE()/DATE() construction matched to the specific Tableau format string."),

    (re.compile(r"no DAX equivalent for Tableau function SCRIPT_\w+\(\)"),
     lambda m: "SCRIPT_* calls an external R/Python integration -- there is no DAX equivalent at "
               "all. Re-implement upstream (a Power Query/dataflow step, or a Power BI R/Python "
               "visual), not as a measure."),

    (re.compile(r"no DAX equivalent for Tableau function REGEXP_\w+\(\)"),
     lambda m: "DAX has no native regex. Rewrite using SEARCH/FIND/SUBSTITUTE for simple "
               "patterns, or precompute a flag/extract column in Power Query (M has richer text "
               "functions) before the model layer."),

    (re.compile(r"no DAX equivalent for Tableau function DATE\(\)"),
     lambda m: "Tableau DATE(expr) converts an expression to a date. Rewrite as "
               "DATEVALUE(expression) for a date-like string, or "
               "DATE(YEAR(expr),MONTH(expr),DAY(expr)) to truncate a datetime."),

    (re.compile(r"no DAX equivalent for Tableau function (\w+)\(\)"),
     lambda m: f"No direct DAX equivalent for {m.group(1)}() -- needs a hand-authored "
               f"restructure; no canned pattern here yet, treat as a one-off human decision."),

    (re.compile(r"STR\(<date>\)"),
     lambda m: 'Rewrite as FORMAT(date, "yyyy-mm-dd") to match Tableau\'s ISO output explicitly '
               "-- DAX's default text coercion is locale-dependent."),

    (re.compile(r"Tableau MID\(s,start\[,len\]\) vs DAX MID"),
     lambda m: "When Tableau omits the length argument, compute it explicitly as "
               "LEN(s) - start + 1 and pass it to DAX MID(s, start, len)."),

    (re.compile(r"Tableau DATEADD\('unit',n,date\) vs DAX DATEADD"),
     lambda m: "Rewrite with DAX's argument order and a keyword unit, e.g. "
               "DATEADD(<date column>, n, MONTH)."),

    (re.compile(r"Tableau DATEDIFF\('unit',a,b\) vs DAX DATEDIFF"),
     lambda m: "Rewrite with DAX's argument order and a keyword unit, e.g. "
               "DATEDIFF(a, b, DAY)."),

    (re.compile(r"Tableau DATETRUNC\('unit',date\) has no single DAX function"),
     lambda m, formula: _unit_fix(formula, _DATETRUNC_UNIT, _DATETRUNC_PATTERNS,
                                   "Tableau DATETRUNC('unit', date)")),

    (re.compile(r"Tableau DATEPART\('unit',date\) has no single DAX function"),
     lambda m, formula: _unit_fix(formula, _DATEPART_UNIT, _DATEPART_PATTERNS,
                                   "Tableau DATEPART('unit', date)")),

    (re.compile(r"which has no aggregate function of its own and was promoted to a DAX "
                r"MEASURE elsewhere"),
     lambda m: "This calc aggregates something that's row-level in Tableau, but this pipeline "
               "already promoted that same referenced calc to a DAX MEASURE elsewhere (forced "
               "by a parameter reference or a downstream aggregate depending on it) -- DAX "
               "cannot re-aggregate a measure with SUM/MIN/MAX/AVG directly. Fix by hand: open "
               "the referenced calc, take its row-level Tableau expression, and inline that "
               "expression directly into an iterator here (e.g. SUMX('Table', <that row-level "
               "expression>)) instead of referencing the promoted measure."),

    (re.compile(r"STR\(<expression containing a date field>\)"),
     lambda m: "The date field is nested inside YEAR()/MONTH()/DAY() (or similar), whose return "
               "value is already a plain INTEGER, not a date -- so no date-style FORMAT() "
               "pattern is needed at all. Rewrite using DAX's own number formatting instead: "
               'FORMAT(MONTH(d), "00") for a zero-padded 2-digit piece, or plain concatenation '
               "with & (DAX auto-converts a number to text) for an unpadded one. Just don't hand "
               "a date-style FORMAT pattern to the integer that YEAR/MONTH/DAY already "
               "produced -- that's the trap this reason exists to prevent."),

    (re.compile(r"Tableau (STARTSWITH|ENDSWITH)\(s,sub\) has no DAX equivalent"),
     lambda m: ("Rewrite as LEFT(s, LEN(sub)) = sub." if m.group(1) == "STARTSWITH"
                else "Rewrite as RIGHT(s, LEN(sub)) = sub.")),

    (re.compile(r"Tableau SPLIT\(s,delim,n\) has no DAX equivalent"),
     lambda m: 'No single DAX function. Rewrite using PATHITEM(SUBSTITUTE(s, delim, "|"), n) '
               "(works when delim is not already '|'), or split upstream in Power Query with "
               "Text.Split."),

    (re.compile(r"Tableau (LTRIM|RTRIM) removes"),
     lambda m: "DAX TRIM() removes both ends; if whitespace on the other side is meaningful, "
               "isolate it manually (e.g. via SUBSTITUTE/MID) rather than swapping in TRIM()."),

    (re.compile(r"circular reference among Tableau calculated fields"),
     lambda m: "This is a genuine cycle in the Tableau source, not a translation gap -- there is "
               "no order to evaluate it in Tableau either. Ask the workbook's author to break "
               "the cycle before this can be translated at all."),

    (re.compile(r"free-entry parameter domain"),
     lambda m: "Convert to a fixed list: find the real column this parameter is compared "
               "against in its own DAX (e.g. inside a SELECTEDVALUE/IF that uses it), get the "
               "DISTINCT real values that column actually holds in the wired-up data source, "
               "and write those as the parameter table's literal rows -- "
               'Table.FromRows({{"Value1"}, {"Value2"}, ...}, type table [Value = text]) -- '
               "instead of a single blank/placeholder row. Also update the parameter's own "
               "SELECTEDVALUE(...) default (its second argument) from \"\" to whichever of "
               "those real values is most common, so the measure has a sane default before the "
               "user picks a slicer value. If no real data is wired up yet, this has to wait "
               "until it is -- there is no source to sample distinct values from.",
     ),

    (re.compile(r"extracted as a single table", re.IGNORECASE),
     lambda m: "This join was collapsed into one flattened table because the build only had a "
               "stub (empty) partition to work with, or the joined tables were all small "
               "reference/lookup tables. Once real data is wired in, re-check: if EITHER side "
               "of this join is genuinely one-to-many against another table already flattened "
               "into the same PBI table (e.g. a transactions/ratings/events table with many "
               "rows per parent row), staying flattened will cross-multiply every one-to-many "
               "hop into a combinatorial row explosion and inflate every SUM/COUNT measure -- "
               "this is REQUIRED to fix, not optional. The fix: split the flattened table back "
               "into one Power BI table per physical Tableau object on each side of the join, "
               "each with its own partition (native connector or Excel/CSV import matching that "
               "object's own sheet/table), and declare this join as a real entry in "
               "relationships.tmdl (fromColumn/toColumn, correct one/many cardinality) instead "
               "of a joined-in column. If every table involved is small and strictly a "
               "one-to-one or reference lookup (a handful of rows, no fan-out risk), staying "
               "flattened is fine and no action is needed -- use judgement based on the real "
               "row counts, not the join type alone."),

    (re.compile(r"relate with no matching fields.*cannot join on a constant"),
     lambda m: "No fix without a real key. Either find a genuine matching field in the source "
               "(check with the workbook author) or leave the tables unrelated and combine what "
               "is needed via a calculated/union table instead of a relationship."),

    (re.compile(r"group source field '.*?' did not resolve"),
     lambda m: "The group's source field was renamed/removed in Tableau -- recreate the group "
               "against the current field name."),

    (re.compile(r"bin needs a resolvable numeric source and size"),
     lambda m: "Verify the bin's source field is numeric and its size is a literal, not a "
               "calculated/dynamic value; replace a dynamic size with a fixed one."),

    (re.compile(r"drill-path level\(s\) resolved to no emitted column"),
     lambda m: "The missing hierarchy levels are themselves declined calcs elsewhere in this "
               "workbook -- fixing those recovers this hierarchy level for free."),

    (re.compile(r"zone type '.*?' has no mapped Power BI visual"),
     lambda m: "No Power BI visual type matches this Tableau zone kind. Pick the closest native "
               "visual by hand in Desktop, or check AppSource for a custom visual for this chart "
               "type."),

    (re.compile(r"no rule fired for mark"),
     lambda m: "Open the worksheet in Tableau to see what it actually renders for this shelf "
               "combination, then manually pick the matching Power BI visual in Desktop -- a "
               "safe `tableEx` placeholder was emitted so nothing is silently wrong."),

    (re.compile(r"filter action maps to|URL action has no Power BI counterpart"),
     lambda m: "Wire the equivalent behavior by hand: Format > Edit interactions for a "
               "filter/highlight action; a URL action needs a button or image with a \"Web URL\" "
               "action, since Power BI has no native URL-action equivalent."),

    (re.compile(r"closest counterpart is a Power BI bookmark"),
     lambda m: "Manually create a Bookmark in Desktop capturing the same state, and wire a "
               "button/action to apply it."),

    (re.compile(r"a `\.hyper` file cannot be read by Power Query"),
     lambda m: "Reconnect this source directly (Import mode) and recreate its refresh schedule "
               "and extract filters as Power Query filter steps or a dataset refresh policy in "
               "the Power BI service -- the .hyper extract itself cannot be read."),

    (re.compile(r"would become a Power Query step or a model-level filter"),
     lambda m: "Implement as a Power Query filter step (Table.SelectRows) or a fixed row-level "
               "filter on the semantic model, matching the datasource filter's condition."),

    (re.compile(r"plumbing, not data"),
     lambda m: "No action needed -- this is a Tableau-internal bookkeeping field with no data "
               "meaning."),

    (re.compile(r"renames a displayed value without changing the data"),
     lambda m: "Add a calculated column mapping value -> alias (small cardinality) or a Power "
               "Query Table.ReplaceValue step, matching the alias pairs from the Tableau source."),

    (re.compile(r"Power Query has no connect-time hook"),
     lambda m: "No direct fix inside Power BI. If a connect-time transform is required, run it "
               "upstream (a scheduled job or dataflow) before Power BI ever reads the source."),

    (re.compile(r"image zone references a file not present"),
     lambda m: "Locate the missing image asset and add it to the report's StaticResources "
               "manually."),

    (re.compile(r"unexpected character|unexpected token|unexpected end of expression|"
                r"trailing tokens"),
     lambda m: "Likely a malformed or corrupted formula in the Tableau source -- open it in "
               "Tableau's calculation editor and re-save; the parser declines rather than "
               "guessing at broken syntax."),
]

GENERIC_FIX = ("No canned fix yet for this reason -- treat as a one-off human decision and "
               "consider adding a rule to FIX_RULES in validate_mapping.py if this pattern "
               "recurs.")


def suggest_fix(reason, formula=""):
    for pat, fn in FIX_RULES:
        m = pat.search(reason or "")
        if m:
            try:
                try:
                    return fn(m, formula)
                except TypeError:
                    return fn(m)
            except Exception:
                return GENERIC_FIX
    return GENERIC_FIX


# ==================================================================== loading
def load_workbook_artifacts(wb_dir):
    work = os.path.join(wb_dir, "_work")
    with open(os.path.join(work, "ir.json"), encoding="utf-8") as fh:
        ir = json.load(fh)
    with open(os.path.join(work, "name_authority.json"), encoding="utf-8") as fh:
        na = json.load(fh)
    residue_path = os.path.join(work, "residue.json")
    if os.path.exists(residue_path):
        with open(residue_path, encoding="utf-8") as fh:
            residue = json.load(fh)
    else:
        residue = {"calc_residue": [], "visual_undecided": [], "zone_gaps": []}
    return ir, na, residue


_TABLE_HEADER = re.compile(r"^table (?:'((?:[^']|'')+)'|(\S+))")
# TMDL only quotes an identifier that needs it (names.py's own `tmdl_ident()` rule); a simple
# name like `Value` or `Channel` is written bare -- `column Value`, not `column 'Value'`. Every
# one of these regexes must accept both forms, matching `_TABLE_HEADER` above, or every unquoted
# name is invisible to the dangling-reference check below: found 124 false positives from this
# exact gap during development, ALL unquoted names like the parameter table's own `Value` column.
_MEASURE_LINE = re.compile(r"^\tmeasure (?:'((?:[^']|'')+)'|(\S+))\s*=\s*(.*)$")
_COLUMN_LINE = re.compile(r"^\tcolumn (?:'((?:[^']|'')+)'|(\S+))\s*=\s*(.*)$")
_COLUMN_BARE_LINE = re.compile(r"^\tcolumn (?:'((?:[^']|'')+)'|(\S+))\s*$")


def load_dax_by_table(wb_dir, model_dirname):
    """Read the emitted TMDL and return (dax_by_table, names_by_table, all_measure_names):

    - dax_by_table[table][name]  -- DAX text, for CALCULATED members only (measures + calc
      columns); used for VALIDATE.dax's query text.
    - names_by_table[table]      -- every real column/measure name on that table, calculated
      or plain physical (`column 'X'` with no expression has none to capture but the name
      still needs to exist in this universe for the dangling-reference check below).
    - all_measure_names          -- every measure name, model-wide -- measures are unique
      across the whole model (the Binding Ledger's own guarantee) and referenced unqualified.

    Most real corpus expressions stay on one physical line, so this reads the measure/column
    line's own tail; a genuinely multi-line one (real example: a group-of-a-group's SWITCH()
    wraps once its branch count grows) degrades to a placeholder string for `dax_by_table`
    rather than a wrong parse. The header line -- `column 'Name' =` with nothing after the `=`
    -- must still be MATCHED even though it captures no inline expression, or the name never
    reaches `names_by_table` and the dangling-reference check below flags a perfectly real,
    correctly-emitted column as missing. A too-strict "= " (must be followed by a space) instead
    of "=\\s*" produced exactly that false positive during development, on this exact case.
    """
    dax_by_table, names_by_table, all_measure_names = {}, {}, set()
    tdir = os.path.join(wb_dir, model_dirname, "definition", "tables")
    if not os.path.isdir(tdir):
        return dax_by_table, names_by_table, all_measure_names
    for fn in sorted(os.listdir(tdir)):
        if not fn.endswith(".tmdl"):
            continue
        with open(os.path.join(tdir, fn), encoding="utf-8") as fh:
            lines = fh.read().splitlines()
        if not lines:
            continue
        hm = _TABLE_HEADER.match(lines[0])
        tname = ((hm.group(1) or hm.group(2)).replace("''", "'")) if hm else fn[:-5]
        members, names = {}, set()
        for ln in lines[1:]:
            mm = _MEASURE_LINE.match(ln)
            if mm:
                name = (mm.group(1) or mm.group(2)).replace("''", "'")
                expr = mm.group(3).strip()
                members[name] = expr if expr else "(multi-line expression -- see TMDL directly)"
                names.add(name)
                all_measure_names.add(name)
                continue
            mm = _COLUMN_LINE.match(ln)
            if mm:
                name = (mm.group(1) or mm.group(2)).replace("''", "'")
                expr = mm.group(3).strip()
                members[name] = expr if expr else "(multi-line expression -- see TMDL directly)"
                names.add(name)
                continue
            mm = _COLUMN_BARE_LINE.match(ln)
            if mm:
                names.add((mm.group(1) or mm.group(2)).replace("''", "'"))
        dax_by_table[tname] = members
        names_by_table[tname] = names
    return dax_by_table, names_by_table, all_measure_names


def _na_key(caption, tableau_name):
    ident = (tableau_name or "").strip("[]")
    return f"{caption} | {ident}"


# ================================================= identifier dereferencing
# Tableau formulas reference other objects by opaque internal identity, not by the caption a
# Tableau user actually sees: `[Parameters].[Parameter 15]` for a parameter, or
# `[Visa Countryregion C (copy)_739153350486216704]` for a field Tableau auto-suffixed with a
# generated id (duplicated/"(copy)" fields get this). Both identities, and their real captions,
# already exist in ir.json -- this is a lookup, not an inference, so it stays fail-closed: an
# identity this workbook does not define is left exactly as written, never guessed at.
def build_identity_map(ir):
    """Raw Tableau identity (tableau_name, brackets stripped) -> {'caption', 'kind', 'table'}.

    Parameters are indexed first so a same-named field can never shadow a parameter's caption.
    `table` is only meaningful for kind == 'field' (which table's Data pane the field lives in);
    parameters are workbook-global in Tableau, so they carry no table."""
    identity_map = {}
    for p in ir.get("parameters", []):
        ident = (p.get("tableau_name") or "").strip("[]")
        if ident and p.get("name"):
            identity_map[ident] = {"caption": p["name"], "kind": "parameter", "table": ""}
    for t in ir.get("tables", []):
        table_caption = t.get("caption") or t.get("tableau_name")
        for coll in (t.get("calcs", []), t.get("columns", [])):
            for c in coll:
                ident = (c.get("tableau_name") or "").strip("[]")
                if ident and c.get("name") and ident not in identity_map:
                    identity_map[ident] = {"caption": c["name"], "kind": "field",
                                            "table": table_caption}
    return identity_map


# Qualified parameter form must be tried before the bare-bracket alternative, or the bare form
# would match just the literal "[Parameters]" segment as if it were a field of its own.
_FORMULA_IDENT = re.compile(r"\[Parameters\]\.\[([^\]]+)\]|\[([^\]]+)\]")


def dereference_formula(text, identity_map, usage=None):
    """Rewrite every `[Parameters].[Parameter N]` and bare `[<opaque identity>]` token in a
    Tableau formula to `[<real caption>]`, e.g. `[Parameters].[Parameter 15]` -> `[Region -
    Parameter]`. One pass over the ORIGINAL text (re.sub never rescans its own replacement), so a
    caption that happens to look like a bracket token is never re-processed. `usage`, if given, is
    a Counter keyed by raw identity, incremented once per substitution -- feeds the Identifier Map
    sheet. An identity absent from `identity_map` is left untouched."""
    if not text:
        return text

    def _sub(m):
        ident = m.group(1) if m.group(1) is not None else m.group(2)
        info = identity_map.get(ident)
        if not info:
            return m.group(0)
        if usage is not None:
            usage[ident] += 1
        return f"[{info['caption']}]"

    return _FORMULA_IDENT.sub(_sub, text)


# DAX's `'Table'[Column]` is already a final, real name -- never a Tableau internal identity --
# so it must be skipped whole (group(1) branch) rather than have its `[Column]` half treated as a
# bare bracket token. `[Parameters].[X]` is included defensively for symmetry even though the
# emitted DAX in this pipeline is always caption-named already (tmdl.py never writes a raw
# `Parameter N` token); if that ever changes, this still resolves it rather than passing it
# through as noise.
_DAX_IDENT = re.compile(r"'((?:[^']|'')+)'\[([^\]]+)\]|\[Parameters\]\.\[([^\]]+)\]|\[([^\]]+)\]")


def dereference_dax(text, identity_map, usage=None):
    """The DAX-side twin of `dereference_formula` -- same identity lookup, applied only to bare
    `[Name]` measure-style tokens; any `'Table'[Column]` qualified reference is left completely
    untouched (it is already a real column/measure name, resolved and validated elsewhere in this
    script's own dangling-reference check)."""
    if not text:
        return text

    def _sub(m):
        if m.group(1) is not None:
            return m.group(0)
        ident = m.group(3) if m.group(3) is not None else m.group(4)
        info = identity_map.get(ident)
        if not info:
            return m.group(0)
        if usage is not None:
            usage[ident] += 1
        return f"[{info['caption']}]"

    return _DAX_IDENT.sub(_sub, text)


_NEWLINES = re.compile(r"\r\n|\r|\n")


def _flatten(text):
    """Collapse embedded line breaks to a single space. A real Tableau multi-line calc (e.g.
    `IF ... THEN TRUE\\r\\nELSE FALSE end`) written as-is into an Excel cell renders as
    `TRUEELSE` with no separator at all -- Excel's cell renderer drops a bare `\\r`/`\\n` rather
    than showing a line break, so the words visually fuse into one. Applied once, to the
    *_readable fields only, so every place this report shows a formula/DAX -- Excel cell or
    VALIDATE.dax comment alike -- is on a single, correctly spaced line."""
    if not text:
        return text
    return _NEWLINES.sub(" ", text)


# ============================================================ static DAX checks
# Things this script can catch WITHOUT a DAX engine -- a "mapped" row is not the same claim as
# "will not error in Desktop". Found the hard way: 145 emitted measures/columns across the
# corpus referenced a parameter table's raw [Value] column instead of its SELECTEDVALUE measure,
# and every one of them errored in Desktop with "a single value for column 'Value'... cannot be
# determined" -- DAX cannot implicitly reduce a multi-row column to a scalar. Every `Param <X>`
# table is this pipeline's OWN synthetic construct (names.py's `parameter()` convention), so this
# pattern is unambiguous, not a heuristic with false positives: any occurrence outside
# SELECTEDVALUE(...) is always wrong.
_PARAM_VALUE_REF = re.compile(r"'(Param [^']*)'\[Value\]")
_QUALIFIED_REF = re.compile(r"'((?:[^']|'')+)'\[([^\]]+)\]")
_BRACKET_TOKEN = re.compile(r"\[([^\]]+)\]")


def find_case_collisions(names_by_table):
    """Per table, any two names that differ only by case.

    TOM's own name collection ignores case (Windows/SQL Server convention) -- a real corpus
    workbook had a physical column `Bic Compliance Check` and a calculated-field copy of it
    captioned `BIC Compliance Check`, and Desktop refused to open the file AT ALL: "Item 'BIC
    Compliance Check' already exists in the collection." Worse than one wrong visual -- a hard
    load failure. The root cause (`names.py`'s collision check compared case-sensitively) is
    fixed, but this check would catch the whole class again if it ever regressed: it reads the
    same real names this script already parses out of the TMDL for the dangling-reference check,
    just asking a different question of them.
    """
    collisions = []
    for table, names in (names_by_table or {}).items():
        seen = {}
        for name in sorted(names):
            folded = name.casefold()
            if folded in seen:
                collisions.append((table, seen[folded], name))
            else:
                seen[folded] = name
    return collisions


def _find_dangling_refs(dax, names_by_table, all_measure_names):
    """Every `'Table'[Column]` and bare `[Measure]` token in `dax` that does not resolve to a
    real name in the emitted TMDL.

    A "mapped" row means the row's OWN name was written; it says nothing about whether the
    THINGS IT REFERENCES were. Found the hard way: a Tableau group built on top of another
    group whose own source never resolved got emitted anyway, with a plausible-looking
    `'Table'[Column]` reference to a column that was never written -- valid-looking DAX that
    fails to load in Desktop ("Column '...' cannot be found"). The root cause (a cascade gap in
    `tmdl.py`'s group/bin resolution) is fixed, but this check catches the whole class: any
    dangling reference, from any cause, present or future.
    """
    if not dax:
        return []
    dangling, remainder = [], dax
    for m in _QUALIFIED_REF.finditer(dax):
        table, col = m.group(1), m.group(2)
        if col not in names_by_table.get(table, ()):
            dangling.append(f"'{table}'[{col}]")
        remainder = remainder.replace(m.group(0), "", 1)
    for m in _BRACKET_TOKEN.finditer(remainder):
        if m.group(1) not in all_measure_names:
            dangling.append(f"[{m.group(1)}]")
    return dangling


def static_dax_warning(dax, names_by_table=None, all_measure_names=None):
    if not dax:
        return None
    msgs = []
    for m in _PARAM_VALUE_REF.finditer(dax):
        if dax[:m.start()].rstrip().endswith(("SELECTEDVALUE(", "SELECTEDVALUE (")):
            continue
        msgs.append(f"References {m.group(1)!r}'s raw [Value] column directly, not through its "
                    f"SELECTEDVALUE measure -- this errors in Desktop with \"a single value for "
                    f"column 'Value'... cannot be determined\" the moment that parameter has "
                    f"more than one row. Should reference the parameter's Selected measure "
                    f"instead.")
        break
    if names_by_table is not None and all_measure_names is not None:
        dangling = _find_dangling_refs(dax, names_by_table, all_measure_names)
        if dangling:
            msgs.append(f"References {', '.join(dangling)}, which {'is' if len(dangling) == 1 else 'are'} "
                        f"not a real column/measure in the emitted model -- this errors in Desktop "
                        f"with \"Column/Measure '...' cannot be found or may not be used in this "
                        f"expression\". Usually means this depends on something that was itself "
                        f"declined; check the Gaps sheet for that dependency.")
    return " ".join(msgs) if msgs else None


# =============================================================== the mapping
def build_mapping(ir, na, residue, dax_by_table, names_by_table=None, all_measure_names=None):
    identity_map = build_identity_map(ir)
    # Tracked separately, not as one merged count: the same identity can be dereferenced inside a
    # Tableau formula (or a row's own `tableau_identity`) and, defensively, inside emitted DAX --
    # two different sheets need to answer "where did THIS side's dereferencing actually happen",
    # not a total that hides which side it came from.
    usage_tableau = collections.Counter()
    usage_pbi = collections.Counter()
    ir_param_by_caption = {p.get("name"): p for p in ir.get("parameters", [])}

    measures = na.get("measures", {})
    columns = na.get("columns", {})
    tables_map = na.get("tables", {})
    # `measures`/`columns` are ALLOCATION -- every calc gets a name before the formula parser
    # gets a chance to decline it (names.py's own documented behavior). `emitted` is the real
    # write. Checking allocation alone misclassified 80 of 98 real gaps as "mapped" in one
    # corpus workbook during development -- an allocated-but-declined name is exactly the
    # dangling reference names.py exists to prevent, and this report must not repeat that bug.
    emitted = set(na.get("emitted", []))

    calc_residue_by_key = {}
    for item in residue.get("calc_residue", []):
        calc_residue_by_key.setdefault((item["table"], item["subject"]), item)

    mapped, gaps, unaccounted = [], [], []
    consumed_residue = set()

    for t in ir.get("tables", []):
        caption = t.get("caption") or t.get("tableau_name")
        pbi_table = tables_map.get(t.get("tableau_name"), caption)

        for calc in t.get("calcs", []):
            # Allocation/emission are keyed by the RESOLVED PBI table name, not the raw Tableau
            # caption -- `sanitize()` strips a trailing parenthetical (e.g. "Foo (DB)" -> "Foo")
            # and collisions rename tables, so caption != pbi_table whenever either happens. Using
            # `caption` here silently misclassified 40 real, translated calcs as "unaccounted" in
            # one corpus workbook during development.
            key = _na_key(pbi_table, calc.get("tableau_name"))
            pbi_name, kind = None, None
            if key in measures and f"{pbi_table} | {measures[key]}" in emitted:
                pbi_name, kind = measures[key], "measure"
            elif key in columns and f"{pbi_table} | {columns[key]}" in emitted:
                pbi_name, kind = columns[key], "column"

            formula = calc.get("formula") or ""
            tableau_identity = calc.get("tableau_name") or ""
            row = {"tableau_table": caption, "tableau_field": calc["name"],
                   "tableau_identity": tableau_identity,
                   "tableau_identity_readable": _flatten(dereference_formula(
                       tableau_identity, identity_map, usage_tableau)),
                   "formula": formula,
                   "formula_readable": _flatten(dereference_formula(
                       formula, identity_map, usage_tableau)),
                   "tableau_kind": calc.get("kind")}

            if pbi_name:
                row.update(pbi_table=pbi_table, pbi_name=pbi_name, kind=kind,
                           dax=dax_by_table.get(pbi_table, {}).get(
                               pbi_name, "(not found in TMDL -- name mismatch, investigate)"))
                mapped.append(row)
                # Mark consumed even though mapped, not just when it becomes a gap -- the
                # leftover residue pass below re-walks residue.json directly and, without this,
                # double-reports anything that started in residue.json but was LATER promoted
                # (e.g. by promote_dax_candidates.py) as both mapped AND still a gap. Caught by
                # promoting Prevent Dashboard's 29 authored candidates: every one of them kept
                # showing up in the Gaps sheet even after correctly appearing in Measures/Columns,
                # because only the gap-outcome branch used to add to this set.
                consumed_residue.add((pbi_table, calc["name"]))
                continue

            rkey = (pbi_table, calc["name"])
            ritem = calc_residue_by_key.get(rkey)
            if ritem:
                consumed_residue.add(rkey)
                row.update(tier=ritem["tier"], reason=ritem["reason"],
                           suggested_fix=suggest_fix(ritem["reason"], row["formula"]))
                gaps.append(row)
            else:
                unaccounted.append(row)

        # plain schema columns (not user-authored calcs) -- row 9's asset
        for col in t.get("columns", []):
            key = _na_key(pbi_table, col.get("tableau_name"))
            if key in columns and f"{pbi_table} | {columns[key]}" in emitted:
                pbi_name = columns[key]
                col_identity = col.get("tableau_name") or ""
                mapped.append({
                    "tableau_table": caption, "tableau_field": col["name"],
                    "tableau_identity": col_identity,
                    "tableau_identity_readable": _flatten(dereference_formula(
                        col_identity, identity_map, usage_tableau)),
                    "formula": "", "formula_readable": "",
                    "tableau_kind": "column", "pbi_table": pbi_table, "pbi_name": pbi_name,
                    "kind": "schema-column", "dax": None,
                })
            # A schema column absent from `columns` with no calc wrapper has no residue path --
            # the model-load preflight would already have failed the build over a missing
            # partition source, so it is not a gap this report needs to explain.

    # parameters -- na['parameters'] is already the authoritative Tableau->PBI mapping (S2's own
    # convention record), no cross-reference needed. `tableau_name` here is actually the
    # parameter's CAPTION (na's own key), not its raw `[Parameter N]` identity -- that raw
    # identity, and the parameter's real domain, come from ir.json's own parameter record,
    # matched by that same caption.
    for tableau_name, rec in na.get("parameters", {}).items():
        ir_p = ir_param_by_caption.get(tableau_name)
        raw_identity = (ir_p.get("tableau_name") or "") if ir_p else ""
        members = ir_p.get("members") if ir_p else None
        domain_values = (", ".join(str(m) for m in members) if members
                          else (str(ir_p["range"]) if ir_p and ir_p.get("range") else ""))
        mapped.append({
            "tableau_table": "Parameters", "tableau_field": tableau_name,
            # already the caption -- na's own key -- so this column has nothing left to
            # dereference; shown as-is for the same reason `formula_readable` is a no-op here.
            "tableau_identity": tableau_name, "tableau_identity_readable": tableau_name,
            "raw_identity": raw_identity, "domain_values": domain_values,
            "formula": "", "formula_readable": "", "tableau_kind": "parameter",
            "pbi_table": rec["table"], "pbi_name": rec["measure"], "kind": "parameter",
            "dax": dax_by_table.get(rec["table"], {}).get(rec["measure"]),
        })

    # Every residue item not claimed by a calc above is a decline from a different asset class
    # entirely -- groups, bins, hierarchies (all tagged tier "D" in RESIDUE.md when the
    # deterministic rule can't resolve its source). Out of this function's per-calc matching by
    # construction, but still real declined objects with a real reason, so still worth a
    # suggested fix -- included directly rather than left for the "gaps == calc residue" count to
    # quietly under-report.
    for item in residue.get("calc_residue", []):
        rkey = (item["table"], item["subject"])
        if rkey in consumed_residue:
            continue
        formula = item.get("formula") or ""
        gaps.append({
            "tableau_table": item["table"], "tableau_field": item["subject"],
            "tableau_identity": "", "tableau_identity_readable": "", "formula": formula,
            "formula_readable": _flatten(dereference_formula(formula, identity_map, usage_tableau)),
            "tableau_kind": "other", "tier": item["tier"], "reason": item["reason"],
            "suggested_fix": suggest_fix(item["reason"], formula),
        })

    for row in mapped:
        row["dax_warning"] = static_dax_warning(row.get("dax"), names_by_table, all_measure_names)
        row["dax_readable"] = _flatten(dereference_dax(row.get("dax"), identity_map, usage_pbi))

    def _identifier_rows(usage):
        rows = [
            {"identifier": f"[{ident}]", "dereferenced": f"[{info['caption']}]",
             "kind": info["kind"], "table": info.get("table", ""), "occurrences": n}
            for ident, info in identity_map.items()
            for n in [usage.get(ident, 0)] if n
        ]
        rows.sort(key=lambda r: (r["kind"], r["identifier"]))
        return rows

    identifier_rows_tableau = _identifier_rows(usage_tableau)
    identifier_rows_pbi = _identifier_rows(usage_pbi)

    return mapped, gaps, unaccounted, identifier_rows_tableau, identifier_rows_pbi


# ==================================================================== VALIDATE.dax
def _dax_query_for(row):
    if row["kind"] in ("measure", "parameter"):
        return f'EVALUATE\n{{ [{row["pbi_name"]}] }}'
    if row["kind"] in ("column", "schema-column"):
        t = row["pbi_table"].replace("'", "''")
        return (f"EVALUATE\nTOPN(10, SELECTCOLUMNS('{t}', "
                f"\"Value\", '{t}'[{row['pbi_name']}]))")
    return None


def build_validate_dax(workbook_name, mapped, gaps):
    lines = [
        f"-- Validation script for {workbook_name}",
        "-- Generated by validate-tableau-pbip-mapping. Open in Power BI Desktop's DAX query",
        "-- view, paste a query block, run it, and compare the result against what the Tableau",
        "-- source formula implies. Measure/parameter names are unique model-wide (the Binding",
        "-- Ledger's own guarantee), so an unqualified [Name] always resolves.",
        "--",
        "-- ==================== MAPPED: measures, columns, parameters ====================",
        "",
    ]
    for r in sorted(mapped, key=lambda r: (r["pbi_table"], r["pbi_name"])):
        q = _dax_query_for(r)
        if not q:
            continue
        lines.append(f"-- Tableau: {r['tableau_table']} :: {r['tableau_field']} "
                     f"({r.get('tableau_identity_readable') or r['tableau_identity']})")
        if r["formula"]:
            readable = r.get("formula_readable") or r["formula"]
            lines.append(f"-- Original formula: {readable.replace(chr(10), ' ').replace(chr(13), '')}")
        lines.append(f"-- Power BI: '{r['pbi_table']}'[{r['pbi_name']}]  ({r['kind']})")
        if r.get("dax_warning"):
            lines.append(f"-- STATIC CHECK WARNING: {r['dax_warning']}")
        lines.append(q)
        lines.append("")

    lines += ["-- ==================== GAPS: not translated, with a suggested fix "
              "====================",
              "-- No query to run for these -- each block is what to do by hand.", ""]
    for g in sorted(gaps, key=lambda g: (g["tableau_table"], g["tableau_field"])):
        lines.append(f"-- GAP [{g['tier']}] {g['tableau_table']} :: {g['tableau_field']}")
        if g["formula"]:
            readable = g.get("formula_readable") or g["formula"]
            lines.append(f"-- Tableau formula: {readable.replace(chr(10), ' ').replace(chr(13), '')}")
        lines.append(f"-- Reason: {g['reason']}")
        wrapped = textwrap.wrap(g["suggested_fix"], width=96) or [""]
        lines.append(f"-- Suggested fix: {wrapped[0]}")
        for w in wrapped[1:]:
            lines.append(f"--                {w}")
        lines.append("")
    return "\n".join(lines) + "\n"


# ==================================================================== Excel
_HEADER_FILL = None
_HEADER_FONT = None


def _style_header(ws):
    from openpyxl.styles import Font, PatternFill
    fill = PatternFill("solid", fgColor="2F5496")
    font = Font(bold=True, color="FFFFFF")
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
    ws.freeze_panes = "A2"


def _autosize(ws):
    for col in ws.columns:
        length = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws.column_dimensions[col[0].column_letter].width = max(10, min(60, length + 2))


def _sheet(wb, title, headers, rows, keys):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for r in rows:
        ws.append([r.get(k, "") for k in keys])
    _style_header(ws)
    _autosize(ws)
    return ws


def build_xlsx(path, workbook_name, mapped, gaps, unaccounted, collisions=None,
                identifier_rows_tableau=None, identifier_rows_pbi=None):
    if openpyxl is None:
        raise RuntimeError("openpyxl is required for the Excel export (pip install openpyxl)")
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    measures = [r for r in mapped if r["kind"] == "measure"]
    columns = [r for r in mapped if r["kind"] in ("column", "schema-column")]
    params = [r for r in mapped if r["kind"] == "parameter"]
    tier_counts = collections.Counter(g["tier"] for g in gaps)
    warned = [r for r in mapped if r.get("dax_warning")]
    collisions = collisions or []
    identifier_rows_tableau = identifier_rows_tableau or []
    identifier_rows_pbi = identifier_rows_pbi or []

    ws = wb.create_sheet("Summary")
    ws.append(["Workbook", workbook_name])
    ws.append(["Mapped measures", len(measures)])
    ws.append(["Mapped calculated/schema columns", len(columns)])
    ws.append(["Mapped parameters", len(params)])
    ws.append(["Total mapped", len(mapped)])
    ws.append([])
    ws.append(["Gaps by tier", ""])
    for tier, n in sorted(tier_counts.items()):
        ws.append([tier, n])
    ws.append(["Total gaps", len(gaps)])
    ws.append([])
    ws.append(["Unaccounted (in source, not mapped, not in residue -- should be 0)",
               len(unaccounted)])
    ws.append(["Mapped but flagged by a static DAX check (see 'DAX Warnings' sheet)",
               len(warned)])
    ws.append(["Case-insensitive name collisions (see 'Name Collisions' sheet, should be 0)",
               len(collisions)])
    ws.append(["Identifiers dereferenced on the Tableau side (see 'Identifier Map (Tableau)')",
               len(identifier_rows_tableau)])
    ws.append(["Identifiers dereferenced on the Power BI side (see 'Identifier Map (Power BI)')",
               len(identifier_rows_pbi)])
    _autosize(ws)

    _sheet(wb, "Measures",
           ["Tableau table", "Tableau field", "Tableau identity", "Formula", "PBI table",
            "PBI measure", "DAX", "DAX warning"],
           measures,
           ["tableau_table", "tableau_field", "tableau_identity_readable", "formula_readable",
            "pbi_table", "pbi_name", "dax_readable", "dax_warning"])
    _sheet(wb, "Columns",
           ["Tableau table", "Tableau field", "Tableau identity", "Formula", "PBI table",
            "PBI column", "Kind", "DAX", "DAX warning"],
           columns,
           ["tableau_table", "tableau_field", "tableau_identity_readable", "formula_readable",
            "pbi_table", "pbi_name", "kind", "dax_readable", "dax_warning"])
    _sheet(wb, "Parameters",
           ["Tableau parameter", "Raw Tableau identity", "Domain values", "PBI table",
            "PBI measure", "DAX"],
           params,
           ["tableau_field", "raw_identity", "domain_values", "pbi_table", "pbi_name",
            "dax_readable"])
    _sheet(wb, "DAX Warnings",
           ["Tableau table", "Tableau field", "PBI table", "PBI name", "Kind", "DAX",
            "Warning"],
           warned,
           ["tableau_table", "tableau_field", "pbi_table", "pbi_name", "kind", "dax_readable",
            "dax_warning"])
    _sheet(wb, "Gaps",
           ["Tableau table", "Tableau field", "Tableau identity", "Formula", "Tier", "Reason",
            "Suggested fix"],
           gaps,
           ["tableau_table", "tableau_field", "tableau_identity_readable", "formula_readable",
            "tier", "reason", "suggested_fix"])
    _sheet(wb, "Unaccounted",
           ["Tableau table", "Tableau field", "Tableau identity", "Formula"],
           unaccounted,
           ["tableau_table", "tableau_field", "tableau_identity_readable", "formula_readable"])
    _sheet(wb, "Name Collisions",
           ["Table", "Name A", "Name B"],
           [{"table": t, "a": a, "b": b} for t, a, b in collisions],
           ["table", "a", "b"])
    _sheet(wb, "Identifier Map (Tableau)",
           ["Raw identifier", "Dereferenced value", "Kind", "Table (if field)", "Occurrences"],
           identifier_rows_tableau,
           ["identifier", "dereferenced", "kind", "table", "occurrences"])
    _sheet(wb, "Identifier Map (Power BI)",
           ["Raw identifier", "Dereferenced value", "Kind", "Table (if field)", "Occurrences"],
           identifier_rows_pbi,
           ["identifier", "dereferenced", "kind", "table", "occurrences"])

    wb.save(path)


# ==================================================================== rollup
def build_rollup_md(rollup):
    lines = ["# Mapping Validation — corpus rollup", "",
             "One row per workbook. `Unaccounted` and `Collisions` should always be 0 -- "
             "`Unaccounted` nonzero means a Tableau field is neither in the Binding Ledger nor "
             "in residue; `Collisions` nonzero means two objects in one table share a name "
             "case-insensitively, which Desktop refuses to load at all. Both are inconsistencies "
             "in the build itself, not a gap in this report.", "",
             "| Workbook | Mapped | Gaps | Unaccounted | Collisions | Gap tiers |",
             "|---|---|---|---|---|---|"]
    tot_mapped = tot_gaps = tot_unacc = tot_coll = 0
    tot_tiers = collections.Counter()
    for r in rollup:
        tiers = ", ".join(f"{k}:{v}" for k, v in sorted(r["tiers"].items())) or "—"
        lines.append(f"| {r['name']} | {r['mapped']} | {r['gaps']} | {r['unaccounted']} | "
                     f"{r.get('collisions', 0)} | {tiers} |")
        tot_mapped += r["mapped"]
        tot_gaps += r["gaps"]
        tot_unacc += r["unaccounted"]
        tot_coll += r.get("collisions", 0)
        tot_tiers.update(r["tiers"])
    tot_tier_str = ", ".join(f"{k}:{v}" for k, v in sorted(tot_tiers.items())) or "—"
    lines.append(f"| **Corpus total** | **{tot_mapped}** | **{tot_gaps}** | **{tot_unacc}** | "
                 f"**{tot_coll}** | {tot_tier_str} |")
    lines.append("")
    return "\n".join(lines) + "\n"


# ======================================================================= main
def main(argv=None):
    ap = argparse.ArgumentParser(
        description="Validate Tableau -> Power BI mapping from a real build-tableau-pbip output")
    ap.add_argument("build_out", help="build-tableau-pbip --out folder (one subfolder/workbook)")
    ap.add_argument("--out", required=True, help="where to write the mapping report")
    args = ap.parse_args(argv)

    os.makedirs(args.out, exist_ok=True)
    rollup = []
    for name in sorted(os.listdir(args.build_out)):
        wb_dir = os.path.join(args.build_out, name)
        work = os.path.join(wb_dir, "_work")
        if not os.path.isdir(work) or not os.path.exists(os.path.join(work, "ir.json")):
            continue

        ir, na, residue = load_workbook_artifacts(wb_dir)
        model_dirname = next((d for d in os.listdir(wb_dir) if d.endswith(".SemanticModel")),
                              None)
        if model_dirname:
            dax_by_table, names_by_table, all_measure_names = load_dax_by_table(wb_dir, model_dirname)
        else:
            dax_by_table, names_by_table, all_measure_names = {}, {}, set()
        mapped, gaps, unaccounted, identifier_rows_tableau, identifier_rows_pbi = build_mapping(
            ir, na, residue, dax_by_table, names_by_table, all_measure_names)
        collisions = find_case_collisions(names_by_table)

        safe = re.sub(r'[<>:"/\\|?*\[\]]', "_", name).strip(". ") or "Workbook"
        wb_out = os.path.join(args.out, safe)
        os.makedirs(wb_out, exist_ok=True)
        build_xlsx(os.path.join(wb_out, f"{safe}_mapping.xlsx"), name, mapped, gaps, unaccounted,
                   collisions, identifier_rows_tableau, identifier_rows_pbi)
        with open(os.path.join(wb_out, "VALIDATE.dax"), "w", encoding="utf-8", newline="\n") as fh:
            fh.write(build_validate_dax(name, mapped, gaps))

        tier_counts = dict(collections.Counter(g["tier"] for g in gaps))
        rollup.append({"name": name, "mapped": len(mapped), "gaps": len(gaps),
                        "unaccounted": len(unaccounted), "tiers": tier_counts,
                        "collisions": len(collisions)})
        print(f"[ok] {name:50s} mapped={len(mapped):4d} gaps={len(gaps):4d} "
              f"unaccounted={len(unaccounted)} collisions={len(collisions)}  -> {wb_out}")

    with open(os.path.join(args.out, "MAPPING_INVENTORY.md"), "w", encoding="utf-8",
              newline="\n") as fh:
        fh.write(build_rollup_md(rollup))
    print(f"\nWritten to {args.out}")


if __name__ == "__main__":
    if len(sys.argv) == 1:
        # self-test
        assert "COUNTROWS(DISTINCT" in suggest_fix(
            "Tableau COUNTD(<expression>) needs a single table to project into "
            "SELECTCOLUMNS(); the expression references 2 (A, B)")
        assert "Changed Type" in suggest_fix(
            "Tableau MIN([x]) aggregates a field Tableau itself declares datatype 'string'")
        assert "ISBLANK" in suggest_fix("comparison against NULL -- DAX treats BLANK() = 0")
        assert "ALLEXCEPT" in suggest_fix(
            "level-of-detail expression -- needs a CALCULATE/ALLEXCEPT restructure")
        assert "recovers automatically" in suggest_fix(
            "depends on 'X', which was itself declined (tier L2) -- translating that one calc "
            "would also recover this")
        f = suggest_fix("Tableau SUM(<expression>) needs the DAX iterator SUMX, which requires "
                         "ONE table to iterate; the expression references 3 (A, B, C)")
        assert "bridge/star-schema" in f, f
        f0 = suggest_fix("Tableau SUM(<expression>) needs the DAX iterator SUMX, which requires "
                          "ONE table to iterate; the expression references 0 (none)")
        assert "authoring artifact" in f0, f0
        assert suggest_fix("something never seen before") == GENERIC_FIX
        assert _dax_query_for({"kind": "measure", "pbi_name": "Total Sales"}) == \
            "EVALUATE\n{ [Total Sales] }"
        assert "SELECTCOLUMNS" in _dax_query_for(
            {"kind": "column", "pbi_table": "Data", "pbi_name": "Region"})
        assert static_dax_warning(
            'IF(\'Param 1.2 View By\'[Value] = "Volume", 1, 0)') is not None
        assert static_dax_warning(
            "SELECTEDVALUE('Param 1.2 View By'[Value])") is None
        assert static_dax_warning("SUM('T'[x])") is None
        assert static_dax_warning(None) is None
        names_by_table = {"T": {"x", "y"}}
        all_measures = {"Total Sales"}
        assert static_dax_warning("SUM('T'[x])", names_by_table, all_measures) is None
        assert static_dax_warning("SUM('T'[x]) + [Total Sales]", names_by_table,
                                   all_measures) is None
        w = static_dax_warning("SWITCH(TRUE(), 'T'[z] = 1, 2, 3)", names_by_table, all_measures)
        assert w is not None and "'T'[z]" in w, w
        w = static_dax_warning("[Missing Measure] + 1", names_by_table, all_measures)
        assert w is not None and "[Missing Measure]" in w, w
        coll = find_case_collisions({"T": {"Bic Compliance Check", "BIC Compliance Check"},
                                      "U": {"Region", "region"}})
        assert len(coll) == 2, coll
        assert find_case_collisions({"T": {"Region", "Sales"}}) == []

        # identifier dereferencing -- real shapes from a real corpus workbook: a parameter
        # ("Parameter 15" -> caption "Region - Parameter") and an auto-suffixed duplicate field
        # ("...(copy)_739153350486216704" -> caption "Region").
        test_ir = {
            "parameters": [{"name": "Region - Parameter", "tableau_name": "[Parameter 15]"}],
            "tables": [{"caption": "Data", "tableau_name": "Data", "calcs": [
                {"name": "Region",
                 "tableau_name": "[Visa Countryregion C (copy)_739153350486216704]"}],
                "columns": []}],
        }
        imap = build_identity_map(test_ir)
        assert imap["Parameter 15"] == {"caption": "Region - Parameter", "kind": "parameter",
                                         "table": ""}
        assert imap["Visa Countryregion C (copy)_739153350486216704"]["caption"] == "Region"

        usage = collections.Counter()
        out = dereference_formula(
            "[Parameters].[Parameter 15] = [Visa Countryregion C (copy)_739153350486216704] "
            "OR [Parameters].[Parameter 15]= 'Global'", imap, usage)
        assert out == "[Region - Parameter] = [Region] OR [Region - Parameter]= 'Global'", out
        assert usage["Parameter 15"] == 2 and usage["Visa Countryregion C (copy)_739153350486216704"] == 1

        # an identity this workbook doesn't define is left exactly as written -- fail closed.
        assert dereference_formula("[Parameters].[Parameter 99]", imap) == "[Parameters].[Parameter 99]"
        assert dereference_formula(None, imap) is None
        assert dereference_formula("", imap) == ""

        # DAX side: a real 'Table'[Column] reference is never touched even though its bracket
        # half could otherwise look like a bare identity token; a bare [Name] IS still checked.
        imap2 = {"Parameter 14": {"caption": "View (E)", "kind": "parameter", "table": ""}}
        assert dereference_dax("'Param 14'[Value] + [Parameter 14]", imap2) == \
            "'Param 14'[Value] + [View (E)]"
        assert dereference_dax("IF([View (E) Selected] = \"Heat Map\", TRUE(), FALSE())",
                                imap2) == "IF([View (E) Selected] = \"Heat Map\", TRUE(), FALSE())"

        # a real multi-line Tableau calc -- written into an Excel cell as-is, the bare \r\n
        # renders as nothing at all ("TRUE" immediately followed by "ELSE": "TRUEELSE"), not a
        # line break -- so `_flatten` must collapse it to a single space before display.
        assert _flatten('IF [View (E)]="Heat Map" THEN TRUE\r\nELSE FALSE end') == \
            'IF [View (E)]="Heat Map" THEN TRUE ELSE FALSE end'
        assert _flatten("a\nb\rc\r\nd") == "a b c d"
        assert _flatten(None) is None
        assert _flatten("") == ""

        print("validate_mapping.py self-test OK "
              f"({len(FIX_RULES)} fix rules, DAX query templates OK, static DAX check OK, "
              "collision check OK, identifier dereferencing OK, newline flattening OK)")
    else:
        main()
