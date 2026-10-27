#!/usr/bin/env python3
"""
Post-data check: does every Blend Key many-to-many relationship (build-tableau-pbip's own
translation of Tableau data blending, see .claude/skills/build-tableau-pbip/scripts/tmdl.py's
`_relationships()`) actually have any real overlapping key values between its two tables?

Why this exists
----------------
tmdl.py's Tier-1 static check (`_shared_columns`'s formula-shape comparison, `_formula_shape`)
catches a shared blend-key column computed two DIFFERENT ways (e.g. one side a concatenation of
two fields, the other a bare field). It cannot catch two columns that are the SAME shape (e.g.
both an ordinary physical column) but whose real-world VALUE DOMAINS never overlap at all --
nothing in the IR distinguishes that from a legitimately shared dimension; it is only detectable
once real (or realistic-sampled) data actually exists in the project. That is what this check
does, DETECT-AND-REPORT ONLY: it never rewrites relationships.tmdl or any other TMDL file (this
pipeline's fixes are applied by hand in Power BI Desktop, guided by COPILOT_FIX_PROMPTS.md -- see
dummy_data_pipeline_orchestrator.py's own docstring).

Real corpus regression this catches: GCLM_EU_GERMANY_COMMERZBANK_AG_v2's Acquisition<->Activation
Blend Key includes `for_view` -- a plain column, same name, same Power BI dataType on both sides
-- whose real sampled values are {opcert, opcert_dist} on one side and {Channel, MCC, Spend Build
Up, ...} on the other: ZERO overlap, so the many-to-many relationship can never produce a matching
row, and any visual spanning both tables renders empty in Power BI Desktop.

What it checks, per project
----------------------------
For every relationship in relationships.tmdl whose join column is a synthetic composite
`Blend Key (col1, col2, ...)` column (tmdl.py's own naming convention -- a SINGLE shared column
becomes a direct, non-composite relationship with no `Blend Key (...)` wrapper at all, out of
scope for this pass, see `find_blend_key_relationships`'s docstring):

  1. Resolve each side's real Excel-wired data (generate-dummy-data's own
     `<Workbook>_dummy_data.xlsx`, reusing validate-dummy-data's own `resolve_real_excel`/
     `parse_excel_wired` rather than re-deriving them).
  2. Recompute each side's actual Blend Key VALUE SET by replaying the identical `"|"`-join the
     TMDL's own DAX performs, over the constituent columns' real row values.
  3. Report every relationship pair with ZERO intersection.

Run
---
    python3 validate_blend_key_domains.py <build output: project-or-parent-folder> [--out DIR]

Writes BLEND_KEY_DOMAIN_REPORT.json + .md into `<out>/<project name>/` when `--out` is given (a
shared directory across several projects needs the per-project nesting; matches
generate-copilot-fix-prompts/scripts/effort_report.py's own `--out` convention), or into
`<project's parent>/validation/` by default -- the SAME directory validate-pbip-compliance/
validate-dummy-data already write into for this one project. Never inside the project's own
build-tableau-pbip output folder, which stays the untouched baseline.

Always exits 0 -- this is a stats-carrying report, not a pass/fail gate (matching
validate_mapping.py/effort_report.py's own convention, not validate_dummy_data.py's): severity
(`zero_overlap`) is read from the printed `[ok]` line or the JSON report, never the exit code.
Needs `openpyxl` (already a dependency of generate-dummy-data itself).

Known limitation, deliberately accepted: value comparison uses `str(value).strip()`, which is not
DAX-exact for a numeric constituent column (Excel `5.0` vs a DAX-rendered `"5"`). This can only
ever produce a FALSE "zero overlap" on a numeric column -- which routes to human review anyway,
the same outcome as any other finding here -- never a false "safe", the one direction this check
exists to prevent.
"""
import argparse
import json
import re
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:                                                       # pragma: no cover
    openpyxl = None

# ---------------------------------------------------------------- wire up this repo's own tools
#
# Mirrors validate-dummy-data/scripts/validate_dummy_data.py's own path setup exactly: four
# levels up from <repo>/.claude/skills/validate-blend-key-domains/scripts/ is the repo root.
_THIS = Path(__file__).resolve()
_REPO_ROOT = _THIS.parents[4]
_TOOLS_DIR = _REPO_ROOT / "tools"
_VDD_DIR = _REPO_ROOT / ".claude" / "skills" / "validate-dummy-data" / "scripts"

for _p in (_TOOLS_DIR, _VDD_DIR):
    if str(_p) not in sys.path:
        sys.path.insert(0, str(_p))

from pbip_rules import model as pbip_model            # noqa: E402  (path set up above)
from validate_dummy_data import (                      # noqa: E402
    find_projects, parse_excel_wired, resolve_real_excel,
)

_BLEND_KEY_RE = re.compile(r"^Blend Key \(")


def find_blend_key_relationships(proj):
    """Relationships whose join column, on BOTH sides, is a synthetic composite `Blend Key (...)`
    column -- build-tableau-pbip/scripts/tmdl.py's own naming convention for a blend pair with
    MORE than one shared column (`_relationships()`: `key_caption = f"Blend Key ({', '.join(shared)})"`).
    A blend pair with exactly one shared column is keyed directly on that column's own name, with
    no composite wrapper at all -- out of scope for this pass: a single physical column's
    domain-overlap risk is no different from an ordinary star-schema relationship's, and is not
    the shape either of this bug's two confirmed regressions took.
    """
    return [rel for rel in proj.relationships
            if _BLEND_KEY_RE.match(rel.from_column or "") and _BLEND_KEY_RE.match(rel.to_column or "")]


def parse_blend_key_constituents(expression):
    """The ordered [col1, col2, ...] a Blend Key column's own TMDL expression concatenates, from
    `'table'[col1] & "|" & 'table'[col2] & ...` -- a bracket-aware scan (mirrors
    build-tableau-pbip/scripts/names.py's own `dax_ref`, which escapes a literal `]` inside a
    column name as `]]`), not a naive `\\[([^\\]]*)\\]` regex that would mis-split on one. Returns
    `None` (declined, not guessed) if the shape doesn't parse -- e.g. an empty/malformed
    expression, or an unterminated `[`.
    """
    if not expression:
        return None
    cols = []
    i, n = 0, len(expression)
    while i < n:
        if expression[i] != "[":
            i += 1
            continue
        j = i + 1
        buf = []
        while j < n:
            if expression[j] == "]":
                if j + 1 < n and expression[j + 1] == "]":
                    buf.append("]")
                    j += 2
                    continue
                break
            buf.append(expression[j])
            j += 1
        if j >= n:
            return None
        cols.append("".join(buf))
        i = j + 1
    return cols or None


def read_excel_rows(xlsx_path, sheet_name):
    """(header, rows) of `sheet_name` in `xlsx_path` -- header is row 1 (trailing blanks trimmed,
    matching validate_dummy_data.read_excel_header's own convention), rows is every row from row 2
    on, as tuples. Returns (None, None, error-detail) on failure."""
    if openpyxl is None:
        return None, None, "openpyxl is not installed -- cannot open the Excel file"
    try:
        wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    except Exception as exc:                                            # noqa: BLE001
        return None, None, f"could not open {xlsx_path.name}: {exc}"
    if sheet_name not in wb.sheetnames:
        return None, None, (f"sheet {sheet_name!r} not found in {xlsx_path.name} "
                             f"(sheets present: {wb.sheetnames})")
    ws = wb[sheet_name]
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
    header = list(header_row)
    while header and header[-1] is None:
        header.pop()
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    return header, rows, None


_STR_LITERAL_RE = re.compile(r'^\s*"((?:[^"]|"")*)"\s*$')
_NUM_LITERAL_RE = re.compile(r'^\s*-?\d+(?:\.\d+)?\s*$')


def literal_constant(expression):
    """If a column's own TMDL `expression` is a bare DAX literal (a quoted string or a plain
    number) with no `[column]` reference in it at all -- e.g. `column blank = " "`, a common
    Tableau blend-padding idiom seen in real corpus data (a constant field included in a blend
    key purely so at least one common field exists) -- its unquoted value; otherwise `None`.
    A constant contributes the SAME value to every row on both sides of a relationship, so it
    never affects whether the two tables' real data actually overlaps; treating it as a stand-in
    value (rather than declining the whole relationship as unresolved just because generate-
    dummy-data never wires a real Excel column for a calc with no source column at all) is safe
    because it is not guessed -- it is read straight from the model's own DAX text."""
    if not expression or "[" in expression:
        return None
    m = _STR_LITERAL_RE.match(expression)
    if m:
        return m.group(1).replace('""', '"')
    if _NUM_LITERAL_RE.match(expression):
        return expression.strip()
    return None


def build_value_set(header, rows, constituent_columns, constants=None):
    """The set of Blend Key values a table's real data would actually produce -- one `"|"`-joined
    string per row, over `constituent_columns` in the SAME order the DAX itself concatenates them
    (tmdl.py's `_relationships()`: `expr = " & \\"|\\" & ".join(dax_ref(t, c) for c in shared)`).
    `constants` (from `literal_constant`) supplies a fixed value for a constituent name that is
    NOT in `header` at all (a calc with no real Excel column, e.g. a bare-literal blend-padding
    field) -- every row uses that same value for that constituent. Returns (value_set,
    missing_columns) -- `missing_columns` is non-empty (and value_set is None) when a constituent
    name is neither in `header` nor resolvable to a constant, declined rather than guessed at."""
    constants = constants or {}
    missing = [c for c in constituent_columns if c not in header and c not in constants]
    if missing:
        return None, missing
    idx = {c: header.index(c) for c in constituent_columns if c in header}
    values = set()
    for row in rows:
        parts = []
        for c in constituent_columns:
            if c in constants:
                parts.append(constants[c])
                continue
            v = row[idx[c]] if idx[c] < len(row) else None
            parts.append("" if v is None else str(v).strip())
        values.add("|".join(parts))
    return values, []


def _resolve_constants(proj, table_name, constituent_columns):
    """{name: literal value} for every constituent whose column on `table_name` is a bare DAX
    literal (see `literal_constant`) -- used by `check_project` to supply a stand-in value for a
    constituent that has no real Excel column at all, without guessing anything not already
    written in the model's own TMDL."""
    table = proj.tables.get(table_name)
    if table is None:
        return {}
    out = {}
    for name in constituent_columns:
        member = table.columns.get(name)
        if member is None:
            continue
        lit = literal_constant(member.expression)
        if lit is not None:
            out[name] = lit
    return out


def check_project(project_root):
    """One project's full Tier-2 check. Never raises -- every failure mode (unresolved
    relationship, missing Excel file, unparseable expression) is recorded as its own
    `unresolved_relationships` entry instead, matching this pipeline's fail-closed/report-don't-
    guess convention. Returns the report dict (see module docstring for its shape)."""
    proj = pbip_model.load_project(project_root)
    real_xlsx = resolve_real_excel(project_root)
    blend_rels = find_blend_key_relationships(proj)

    # Cache each table's (header, rows) once, even if it participates in several blend
    # relationships (a real corpus workbook does).
    table_cache = {}

    def table_rows(table_name):
        if table_name in table_cache:
            return table_cache[table_name]
        table = proj.tables.get(table_name)
        if table is None:
            table_cache[table_name] = (None, None, f"table {table_name!r} not found in the model")
            return table_cache[table_name]
        source = next((p.source for p in table.partitions
                        if "Excel.Workbook(File.Contents(" in (p.source or "")), "")
        parsed = parse_excel_wired(source) if source else None
        if parsed is None:
            table_cache[table_name] = (
                None, None,
                f"table {table_name!r} is not wired to a real Excel.Workbook(...) partition -- "
                f"no real data to check its Blend Key values against")
            return table_cache[table_name]
        if real_xlsx is None:
            table_cache[table_name] = (
                None, None, "no '<Workbook>_dummy_data.xlsx' file found in this project's own "
                             "folder to read")
            return table_cache[table_name]
        header, rows, err = read_excel_rows(real_xlsx, parsed["sheet"])
        table_cache[table_name] = (header, rows, err)
        return table_cache[table_name]

    relationships_out, unresolved = [], []
    for rel in blend_rels:
        a, a_col, b, b_col = rel.to_table, rel.to_column, rel.from_table, rel.from_column
        member_a = proj.tables[a].columns.get(a_col) if a in proj.tables else None
        member_b = proj.tables[b].columns.get(b_col) if b in proj.tables else None
        expr = (member_a.expression if member_a else "") or (member_b.expression if member_b else "")
        constituents = parse_blend_key_constituents(expr)
        if not constituents:
            unresolved.append({
                "table_a": a, "table_b": b,
                "reason": "the Blend Key column's own TMDL expression did not parse into the "
                          "expected '[c1] & \"|\" & [c2] & ...' shape"})
            continue

        header_a, rows_a, err_a = table_rows(a)
        header_b, rows_b, err_b = table_rows(b)
        if err_a or err_b:
            unresolved.append({"table_a": a, "table_b": b, "reason": err_a or err_b})
            continue

        constants_a = _resolve_constants(proj, a, constituents)
        constants_b = _resolve_constants(proj, b, constituents)
        values_a, missing_a = build_value_set(header_a, rows_a, constituents, constants_a)
        values_b, missing_b = build_value_set(header_b, rows_b, constituents, constants_b)
        if missing_a or missing_b:
            unresolved.append({
                "table_a": a, "table_b": b,
                "reason": f"constituent column(s) not found in the wired Excel header -- "
                          f"{a!r}: {missing_a or 'ok'}, {b!r}: {missing_b or 'ok'}"})
            continue

        overlap = values_a & values_b
        overlap_ratio = (len(overlap) / min(len(values_a), len(values_b))
                          if values_a and values_b else 0.0)
        relationships_out.append({
            "table_a": a, "table_b": b,
            "constituent_columns": constituents,
            "value_count_a": len(values_a), "value_count_b": len(values_b),
            "overlap_count": len(overlap),
            "overlap_ratio": round(overlap_ratio, 4),
            "zero_overlap": len(overlap) == 0,
        })

    return {
        "workbook": project_root.name,
        "checked": len(relationships_out),
        "zero_overlap": sum(1 for r in relationships_out if r["zero_overlap"]),
        "unresolved": len(unresolved),
        "relationships": relationships_out,
        "unresolved_relationships": unresolved,
    }


def render_markdown(report):
    lines = [f"# Blend Key Domain-Overlap Report -- {report['workbook']}", ""]
    if not report["relationships"] and not report["unresolved_relationships"]:
        lines.append("No composite `Blend Key (...)` relationships found in this project.")
        return "\n".join(lines) + "\n"
    if report["zero_overlap"]:
        lines.append(
            f"**{report['zero_overlap']} of {report['checked']} Blend Key relationship(s) have "
            f"ZERO overlapping real values** -- these relationships can never produce a matching "
            f"row; any visual spanning both tables will render empty in Power BI Desktop "
            f"regardless of how much data is loaded.")
        lines.append("")
    lines.append("| Table A | Table B | Constituent columns | Values A | Values B | Overlap | Zero overlap? |")
    lines.append("|---|---|---|---|---|---|---|")
    for r in report["relationships"]:
        flag = "**YES**" if r["zero_overlap"] else "no"
        lines.append(f"| {r['table_a']} | {r['table_b']} | {', '.join(r['constituent_columns'])} "
                      f"| {r['value_count_a']} | {r['value_count_b']} | {r['overlap_count']} "
                      f"({r['overlap_ratio']:.1%}) | {flag} |")
    if report["unresolved_relationships"]:
        lines.append("")
        lines.append("## Unresolved (could not check)")
        lines.append("")
        lines.append("| Table A | Table B | Reason |")
        lines.append("|---|---|---|")
        for u in report["unresolved_relationships"]:
            lines.append(f"| {u['table_a']} | {u['table_b']} | {u['reason']} |")
    return "\n".join(lines) + "\n"


def main():
    ap = argparse.ArgumentParser(
        description="Check every Blend Key relationship's real data for overlapping key values")
    ap.add_argument("input", help="build-tableau-pbip output: project folder, or parent of many")
    ap.add_argument("--out", help="write reports under <out>/<project name>/ instead of in place")
    a = ap.parse_args()

    projects = find_projects(a.input)
    if not projects:
        print(f"[fatal] no PBIP projects found under {a.input}", file=sys.stderr)
        sys.exit(2)

    for project_root in projects:
        report = check_project(project_root)
        out_dir = (Path(a.out).resolve() / project_root.name) if a.out else (project_root.parent / "validation")
        out_dir.mkdir(parents=True, exist_ok=True)
        json_path = out_dir / "BLEND_KEY_DOMAIN_REPORT.json"
        md_path = out_dir / "BLEND_KEY_DOMAIN_REPORT.md"
        json_path.write_text(json.dumps(report, indent=2))
        md_path.write_text(render_markdown(report))

        print(f"[ok] {report['workbook']}  checked={report['checked']} "
              f"zero_overlap={report['zero_overlap']} unresolved={report['unresolved']}  "
              f"-> {json_path}")

    # Always exit 0, deliberately: this is a stats-carrying report, not a pass/fail gate --
    # exactly like validate_mapping.py/effort_report.py, its own severity (`zero_overlap`) is
    # read from the printed [ok] line / JSON, never from the exit code. validator_agent.py's
    # generic call_skill() dispatcher treats ANY non-zero exit as "the skill itself crashed" and
    # discards stdout entirely -- exiting non-zero here (as an earlier version of this script
    # did, matching validate_dummy_data.py's different, genuinely pass/fail convention) silently
    # threw away every real finding the moment one was found, the exact opposite of this
    # script's purpose.
    sys.exit(0)


if __name__ == "__main__":
    main()
